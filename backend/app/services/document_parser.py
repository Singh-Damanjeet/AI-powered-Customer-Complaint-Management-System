"""Safe, in-memory parsing for supported complaint document formats."""

from __future__ import annotations

from email import policy
from email.parser import BytesParser
from html.parser import HTMLParser
from io import BytesIO
from pathlib import Path
import re
from typing import Any


MAX_DOCUMENT_SIZE_BYTES = 10 * 1024 * 1024
SUPPORTED_DOCUMENT_EXTENSIONS = frozenset({".pdf", ".docx", ".txt", ".eml"})


class DocumentParserError(ValueError):
    """Base error for invalid or unreadable complaint documents."""


class DocumentParserConfigurationError(RuntimeError):
    """Raised when an optional format parser is not installed."""


class UnsupportedDocumentTypeError(DocumentParserError):
    """Raised when the filename extension is not supported."""


class EmptyDocumentError(DocumentParserError):
    """Raised when an uploaded file contains no bytes."""


class DocumentTooLargeError(DocumentParserError):
    """Raised when a document exceeds the configured in-memory limit."""


class CorruptDocumentError(DocumentParserError):
    """Raised when a supported document cannot be parsed."""


class NoExtractableTextError(DocumentParserError):
    """Raised when a valid document has no usable textual content."""


# Compatibility aliases keep parser errors discoverable under common names.
UnsupportedFileTypeError = UnsupportedDocumentTypeError
FileTooLargeError = DocumentTooLargeError


_BLOCK_HTML_TAGS = frozenset(
    {
        "address",
        "article",
        "aside",
        "blockquote",
        "br",
        "div",
        "dl",
        "dt",
        "dd",
        "fieldset",
        "figcaption",
        "figure",
        "footer",
        "form",
        "h1",
        "h2",
        "h3",
        "h4",
        "h5",
        "h6",
        "header",
        "hr",
        "li",
        "main",
        "nav",
        "ol",
        "p",
        "pre",
        "section",
        "table",
        "tr",
        "ul",
    }
)


class _ReadableHTMLParser(HTMLParser):
    """Small standard-library HTML-to-text converter for HTML-only emails."""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.parts: list[str] = []
        self._ignored_depth = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        del attrs
        normalized_tag = tag.casefold()
        if normalized_tag in {"script", "style"}:
            self._ignored_depth += 1
            return
        if self._ignored_depth == 0 and normalized_tag in _BLOCK_HTML_TAGS:
            self.parts.append("\n")

    def handle_startendtag(
        self,
        tag: str,
        attrs: list[tuple[str, str | None]],
    ) -> None:
        self.handle_starttag(tag, attrs)
        if tag.casefold() in {"script", "style"}:
            self.handle_endtag(tag)

    def handle_endtag(self, tag: str) -> None:
        normalized_tag = tag.casefold()
        if normalized_tag in {"script", "style"}:
            self._ignored_depth = max(0, self._ignored_depth - 1)
            return
        if self._ignored_depth == 0 and normalized_tag in _BLOCK_HTML_TAGS:
            self.parts.append("\n")

    def handle_data(self, data: str) -> None:
        if self._ignored_depth == 0:
            self.parts.append(data)


def _html_to_text(value: str) -> str:
    parser = _ReadableHTMLParser()
    try:
        parser.feed(value)
        parser.close()
    except Exception as exc:
        raise CorruptDocumentError("The email HTML body could not be read.") from exc
    return "".join(parser.parts)


def _normalize_text(value: str) -> str:
    """Normalize line endings, trailing whitespace, and blank-line runs."""

    normalized = value.replace("\x00", " ")
    normalized = normalized.replace("\r\n", "\n").replace("\r", "\n")
    lines = [line.rstrip() for line in normalized.split("\n")]
    normalized = "\n".join(lines)
    normalized = re.sub(r"\n[ \t]*\n(?:[ \t]*\n)+", "\n\n", normalized)
    return normalized.strip()


def _no_text_error(extension: str) -> NoExtractableTextError:
    if extension == ".pdf":
        return NoExtractableTextError(
            "No extractable text was found in this PDF. OCR is not supported "
            "in this version."
        )
    return NoExtractableTextError("No usable text was found in this document.")


def _decode_email_part(part: Any) -> str:
    try:
        content = part.get_content()
    except Exception:
        payload = part.get_payload(decode=True)
        if payload is None:
            return ""
        charset = part.get_content_charset() or "utf-8"
        try:
            content = payload.decode(charset, errors="replace")
        except (LookupError, UnicodeDecodeError):
            content = payload.decode("utf-8", errors="replace")
    return content if isinstance(content, str) else str(content)


def _parse_email(content: bytes) -> str:
    try:
        message = BytesParser(policy=policy.default).parsebytes(content)
    except Exception as exc:
        raise CorruptDocumentError("The email document could not be parsed.") from exc

    header_lines: list[str] = []
    for label, header_name in (
        ("Subject", "Subject"),
        ("From", "From"),
        ("To", "To"),
        ("Date", "Date"),
    ):
        value = message.get(header_name)
        if value is not None and str(value).strip():
            header_lines.append(f"{label}: {str(value).strip()}")

    plain_parts: list[str] = []
    html_parts: list[str] = []
    parts = message.walk() if message.is_multipart() else [message]
    for part in parts:
        if part.is_multipart():
            continue
        if (
            part.get_content_disposition() == "attachment"
            or part.get_filename() is not None
        ):
            continue
        content_type = part.get_content_type().casefold()
        if content_type == "text/plain":
            plain_parts.append(_decode_email_part(part))
        elif content_type == "text/html":
            html_parts.append(_decode_email_part(part))

    body = "\n".join(plain_parts).strip()
    if not body and html_parts:
        body = _html_to_text("\n".join(html_parts))

    pieces = [*header_lines, body]
    return "\n\n".join(piece for piece in pieces if piece.strip())


def _parse_docx(content: bytes) -> str:
    try:
        from docx import Document
    except ImportError as exc:
        raise DocumentParserConfigurationError(
            "DOCX support requires the python-docx package."
        ) from exc

    try:
        document = Document(BytesIO(content))
        pieces = [paragraph.text for paragraph in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                if any(cells):
                    pieces.append(" | ".join(cells))
        return "\n".join(pieces)
    except Exception as exc:
        raise CorruptDocumentError("The DOCX document could not be parsed.") from exc


def _parse_pdf(content: bytes) -> str:
    try:
        import fitz
    except ImportError as exc:
        raise DocumentParserConfigurationError(
            "PDF support requires the PyMuPDF package."
        ) from exc

    try:
        document = fitz.open(stream=content, filetype="pdf")
        try:
            return "\n".join(page.get_text("text") for page in document)
        finally:
            document.close()
    except Exception as exc:
        raise CorruptDocumentError("The PDF document could not be parsed.") from exc


class DocumentParserService:
    """Validate and parse PDF, DOCX, TXT, and EML files in memory."""

    def __init__(
        self,
        max_size_bytes: int = MAX_DOCUMENT_SIZE_BYTES,
        *,
        max_file_size: int | None = None,
    ) -> None:
        if max_file_size is not None:
            if max_size_bytes != MAX_DOCUMENT_SIZE_BYTES:
                raise ValueError(
                    "Provide either max_size_bytes or max_file_size, not both."
                )
            max_size_bytes = max_file_size
        if not isinstance(max_size_bytes, int) or max_size_bytes <= 0:
            raise ValueError("max_size_bytes must be a positive integer.")
        self.max_size_bytes = max_size_bytes

    async def parse(self, filename: str, content: bytes) -> str:
        """Return normalized text or a clean parser validation error."""

        if not isinstance(filename, str) or not filename.strip():
            raise UnsupportedDocumentTypeError(
                "A filename with a supported extension is required."
            )
        filename_value = filename.strip()
        extension = Path(filename_value).suffix.casefold()
        if extension not in SUPPORTED_DOCUMENT_EXTENSIONS:
            supported = ", ".join(sorted(SUPPORTED_DOCUMENT_EXTENSIONS))
            raise UnsupportedDocumentTypeError(
                f"Unsupported document type. Supported extensions: {supported}."
            )
        if not isinstance(content, (bytes, bytearray, memoryview)):
            raise DocumentParserError("Uploaded document content must be bytes.")

        raw_content = bytes(content)
        if not raw_content:
            raise EmptyDocumentError("Uploaded document is empty.")
        if len(raw_content) > self.max_size_bytes:
            if self.max_size_bytes % (1024 * 1024) == 0:
                limit = f"{self.max_size_bytes // (1024 * 1024)} MB"
            else:
                limit = f"{self.max_size_bytes} bytes"
            raise DocumentTooLargeError(
                f"Uploaded document exceeds the {limit} limit."
            )

        if extension == ".pdf":
            text = _parse_pdf(raw_content)
        elif extension == ".docx":
            text = _parse_docx(raw_content)
        elif extension == ".txt":
            try:
                text = raw_content.decode("utf-8-sig")
            except UnicodeDecodeError as exc:
                raise DocumentParserError(
                    "The text document could not be decoded as UTF-8."
                ) from exc
        else:
            text = _parse_email(raw_content)

        normalized = _normalize_text(text)
        if not normalized:
            raise _no_text_error(extension)
        return normalized


__all__ = [
    "CorruptDocumentError",
    "DocumentParserConfigurationError",
    "DocumentParserError",
    "DocumentParserService",
    "DocumentTooLargeError",
    "EmptyDocumentError",
    "FileTooLargeError",
    "MAX_DOCUMENT_SIZE_BYTES",
    "NoExtractableTextError",
    "SUPPORTED_DOCUMENT_EXTENSIONS",
    "UnsupportedDocumentTypeError",
    "UnsupportedFileTypeError",
]
