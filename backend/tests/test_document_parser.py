"""Tests for safe PDF, DOCX, TXT, and EML parsing."""

import asyncio
from email.message import EmailMessage
from io import BytesIO

import fitz
import pytest
from docx import Document

from app.services.document_parser import (
    CorruptDocumentError,
    DocumentParserService,
    DocumentTooLargeError,
    EmptyDocumentError,
    NoExtractableTextError,
    UnsupportedDocumentTypeError,
)


def run(coroutine):
    return asyncio.run(coroutine)


def text_pdf(text: str | None = None) -> bytes:
    document = fitz.open()
    page = document.new_page()
    if text is not None:
        page.insert_text((72, 72), text)
    content = document.tobytes()
    document.close()
    return content


def text_docx() -> bytes:
    document = Document()
    document.add_paragraph("NovaCure Labs reported a foreign particle.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Batch"
    table.cell(0, 1).text = "AT26007"
    table.cell(1, 0).text = "Product"
    table.cell(1, 1).text = "Atorvastatin Tablets"
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def plain_email() -> bytes:
    message = EmailMessage()
    message["Subject"] = "Product Quality Complaint"
    message["From"] = "customer@example.com"
    message["To"] = "qa@example.com"
    message["Date"] = "Mon, 20 Jul 2026 09:30:00 +0000"
    message.set_content("We observed broken seals on batch PC26018.")
    return message.as_bytes()


def html_only_email() -> bytes:
    message = EmailMessage()
    message["Subject"] = "HTML complaint"
    message["From"] = "customer@example.com"
    message.set_content("<p>Several blister packs were damaged.</p>", subtype="html")
    return message.as_bytes()


def test_pdf_parser_extracts_text() -> None:
    parsed = run(
        DocumentParserService().parse(
            "complaint.PDF",
            text_pdf("Customer: ABC Pharma\nBatch: MT24003"),
        )
    )

    assert "Customer: ABC Pharma" in parsed
    assert "Batch: MT24003" in parsed


def test_empty_pdf_reports_ocr_limitation() -> None:
    with pytest.raises(NoExtractableTextError, match="OCR is not supported"):
        run(DocumentParserService().parse("scan.pdf", text_pdf()))


def test_corrupt_pdf_is_handled_cleanly() -> None:
    with pytest.raises(CorruptDocumentError, match="PDF document could not be parsed"):
        run(DocumentParserService().parse("broken.pdf", b"not a PDF"))


def test_docx_parser_extracts_paragraphs_and_simple_tables() -> None:
    parsed = run(DocumentParserService().parse("foreign_material.docx", text_docx()))

    assert "NovaCure Labs reported a foreign particle." in parsed
    assert "Batch | AT26007" in parsed
    assert "Product | Atorvastatin Tablets" in parsed


def test_txt_parser_decodes_utf8_and_normalizes_blank_lines() -> None:
    parsed = run(
        DocumentParserService().parse(
            "complaint.txt",
            "Line one\r\n\r\n\r\n\r\nLine two – café".encode("utf-8"),
        )
    )

    assert parsed == "Line one\n\nLine two – café"


def test_eml_parser_prefers_plain_text_and_includes_headers() -> None:
    parsed = run(DocumentParserService().parse("complaint.eml", plain_email()))

    assert "Subject: Product Quality Complaint" in parsed
    assert "From: customer@example.com" in parsed
    assert "To: qa@example.com" in parsed
    assert "Date: Mon, 20 Jul 2026 09:30:00 +0000" in parsed
    assert "We observed broken seals on batch PC26018." in parsed


def test_eml_parser_converts_html_only_body() -> None:
    parsed = run(DocumentParserService().parse("complaint.eml", html_only_email()))

    assert "Several blister packs were damaged." in parsed


def test_unsupported_file_is_rejected_before_parsing() -> None:
    with pytest.raises(UnsupportedDocumentTypeError, match="Unsupported document type"):
        run(DocumentParserService().parse("complaint.exe", b"complaint"))


def test_empty_file_is_rejected() -> None:
    with pytest.raises(EmptyDocumentError, match="Uploaded document is empty"):
        run(DocumentParserService().parse("complaint.txt", b""))


def test_file_size_limit_is_configurable() -> None:
    parser = DocumentParserService(max_size_bytes=4)

    with pytest.raises(DocumentTooLargeError, match="limit"):
        run(parser.parse("complaint.txt", b"12345"))
